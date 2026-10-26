"""Aggregations behind the organizer dashboard, revenue reports and exports."""
import csv
import io
from datetime import timedelta
from decimal import Decimal

from django.db.models import Count, DecimalField, F, Q, Sum
from django.db.models.functions import TruncDate
from django.utils import timezone

from apps.checkin.models import CheckIn
from apps.events.models import Event
from apps.registrations.models import Order, OrderItem
from apps.tickets.models import Ticket, TicketType

ZERO = Decimal("0.00")


def organizer_events(user):
    qs = Event.objects.all()
    return qs if user.is_superuser else qs.filter(organizer=user)


def summary(events):
    paid = Order.objects.filter(event__in=events, status=Order.Status.PAID)
    tickets = Ticket.objects.filter(event__in=events).exclude(status=Ticket.Status.CANCELLED)
    return {
        "revenue": paid.aggregate(s=Sum("total"))["s"] or ZERO,
        "orders": paid.count(),
        "tickets_sold": tickets.count(),
        "checked_in": tickets.filter(status=Ticket.Status.USED).count(),
        "pending_orders": Order.objects.filter(
            event__in=events, status=Order.Status.PENDING, expires_at__gt=timezone.now()
        ).count(),
        "refunded": Order.objects.filter(event__in=events, status=Order.Status.REFUNDED).aggregate(s=Sum("total"))["s"]
        or ZERO,
        "events_live": events.filter(status=Event.Status.PUBLISHED, ends_at__gte=timezone.now()).count(),
        "events_total": events.count(),
    }


def daily_sales(events, days=30):
    since = timezone.now() - timedelta(days=days - 1)
    rows = (
        Order.objects.filter(event__in=events, status=Order.Status.PAID, paid_at__gte=since)
        .annotate(day=TruncDate("paid_at"))
        .values("day")
        .annotate(revenue=Sum("total"), orders=Count("id"))
        .order_by("day")
    )
    by_day = {r["day"]: r for r in rows}
    labels, revenue, orders = [], [], []
    start = timezone.localdate() - timedelta(days=days - 1)
    for i in range(days):
        d = start + timedelta(days=i)
        labels.append(d.strftime("%d %b"))
        revenue.append(float(by_day.get(d, {}).get("revenue") or 0))
        orders.append(by_day.get(d, {}).get("orders") or 0)
    return {"labels": labels, "revenue": revenue, "orders": orders}


def ticket_type_breakdown(events):
    rows = []
    for tt in TicketType.objects.filter(event__in=events).select_related("event"):
        rev = OrderItem.objects.filter(ticket_type=tt, order__status=Order.Status.PAID).aggregate(
            s=Sum(F("unit_price") * F("quantity"), output_field=DecimalField())
        )["s"] or ZERO
        rows.append({"event": tt.event.title, "name": tt.name, "price": tt.price, "quantity": tt.quantity,
                     "sold": tt.sold, "revenue": rev, "pct": round(100 * tt.sold / tt.quantity) if tt.quantity else 0})
    return rows


def event_breakdown(events):
    return (
        events.annotate(
            revenue=Sum("orders__total", filter=Q(orders__status=Order.Status.PAID)),
            paid_orders=Count("orders", filter=Q(orders__status=Order.Status.PAID), distinct=True),
        )
        .select_related("venue")
        .order_by("-starts_at")
    )


def checkin_stats(event):
    tickets = Ticket.objects.filter(event=event).exclude(status=Ticket.Status.CANCELLED)
    total = tickets.count()
    inside = tickets.filter(status=Ticket.Status.USED).count()
    return {
        "total": total,
        "checked_in": inside,
        "remaining": total - inside,
        "pct": round(100 * inside / total) if total else 0,
        "rejected": CheckIn.objects.filter(event=event).exclude(result=CheckIn.Result.ACCEPTED).count(),
    }


# ------------------------------------------------------------------ exports
ATTENDEE_HEADERS = ["Ticket", "Attendee", "Email", "Ticket type", "Order", "Status", "Checked in at", "Issued at"]
ORDER_HEADERS = ["Reference", "Event", "Buyer", "Email", "Phone", "Tickets", "Subtotal", "Discount", "Total",
                 "Status", "Discount code", "Created", "Paid"]


def _fmt(dt):
    return timezone.localtime(dt).strftime("%Y-%m-%d %H:%M") if dt else ""


def attendee_rows(event):
    for t in Ticket.objects.filter(event=event).select_related("ticket_type", "order").order_by("attendee_name"):
        yield [t.short_code, t.attendee_name, t.attendee_email, t.ticket_type.name, t.order.reference,
               t.get_status_display(), _fmt(t.checked_in_at), _fmt(t.issued_at)]


def order_rows(events):
    qs = Order.objects.filter(event__in=events).select_related("event", "discount_code").prefetch_related("items")
    for o in qs:
        yield [o.reference, o.event.title, o.buyer_name, o.buyer_email, o.buyer_phone, o.quantity,
               float(o.subtotal), float(o.discount_amount), float(o.total), o.get_status_display(),
               o.discount_code.code if o.discount_code else "", _fmt(o.created_at), _fmt(o.paid_at)]


def to_csv(headers, rows):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(headers, rows, title="Export"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
    for row in rows:
        ws.append(row)
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(c.value or "")) for c in ws[get_column_letter(i)]])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 50)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
